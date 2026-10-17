"""Installed-wheel end-to-end fixture, independent of repository samples."""

from __future__ import annotations

import tempfile
from importlib.resources import files
from pathlib import Path

from pptx import Presentation

from .config import load_config
from .media import write_tone
from .models import Project
from .pipeline import Pipeline, create_job


def smoke(out: Path) -> Project:
    """Exercise FFmpeg, packaging, subtitles, reporting and cache without network/GPU."""
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        cfg_path = root / "smoke.yaml"
        cfg_path.write_text(
            files("slidevoice.configs").joinpath("smoke.yaml").read_text(encoding="utf-8"),
            encoding="utf-8",
        )
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "SlideVoice smoke test"
        slide.notes_slide.notes_text_frame.text = "Xin chào. Đây là âm thanh kiểm thử tổng hợp."
        slides = root / "smoke.pptx"
        deck.save(str(slides))
        voice = root / "tone.wav"
        write_tone(voice, 1)
        job = create_job(slides, voice, out, load_config(cfg_path))
    return Pipeline(job).run()
