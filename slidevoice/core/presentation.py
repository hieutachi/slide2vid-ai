"""Validated lecture plans and PPTX targets, without changing stage contracts."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Literal

from pptx import Presentation
from pydantic import BaseModel, ConfigDict, Field, model_validator

from .cache import write_json
from .errors import InputInvalid
from .models import Project


class PlanModel(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)


class Bounds(PlanModel):
    x: float = Field(ge=0, le=1, allow_inf_nan=False)
    y: float = Field(ge=0, le=1, allow_inf_nan=False)
    width: float = Field(gt=0, le=1, allow_inf_nan=False)
    height: float = Field(gt=0, le=1, allow_inf_nan=False)

    @model_validator(mode="after")
    def inside_slide(self) -> Bounds:
        if self.x + self.width > 1.000001 or self.y + self.height > 1.000001:
            raise ValueError("Target extends outside slide")
        return self


class Target(PlanModel):
    id: str
    text: str
    kind: str
    bounds: Bounds
    parent_id: str | None = None
    geometry_source: Literal["pptx_shape", "powerpoint_text_line"] = "pptx_shape"
    is_text: bool = False


class Evidence(PlanModel):
    source: Literal["slide", "notes"]
    quote: str = Field(min_length=1, max_length=1500)


class Beat(PlanModel):
    text: str = Field(min_length=1, max_length=2000)
    purpose: Literal["orient", "explain", "example", "recap", "transition"]
    evidence: list[Evidence] = Field(min_length=1, max_length=5)
    effect: Literal["none", "highlight", "circle", "zoom"]
    target_id: str | None

    @model_validator(mode="after")
    def valid_target(self) -> Beat:
        if (self.effect == "none") != (self.target_id is None):
            raise ValueError("Effects require a target; none requires null")
        return self


class LecturePlan(PlanModel):
    objective: str = Field(min_length=1, max_length=500)
    beats: list[Beat] = Field(min_length=1, max_length=16)
    takeaway: str = Field(min_length=1, max_length=500)
    warnings: list[str] = Field(max_length=20)


def extract_layout(
    source: Path, text_lines: list[dict[str, Any]] | None = None
) -> dict[str, list[dict[str, Any]]]:
    """Use actual shape geometry, never LLM-invented coordinates.

    Groups are deliberately one coarse target; native animation and per-glyph
    bounds are not available from the exported still image.
    """
    deck = Presentation(str(source))
    if deck.slide_width is None or deck.slide_height is None:
        raise InputInvalid("PPTX slide dimensions are missing.")
    result: dict[str, list[dict[str, Any]]] = {}
    for index, slide in enumerate(deck.slides, 1):
        targets = []
        for shape in slide.shapes:
            left = float(shape.left) / deck.slide_width
            top = float(shape.top) / deck.slide_height
            x = max(0.0, left)
            y = max(0.0, top)
            w = min(left + float(shape.width) / deck.slide_width, 1.0) - x
            h = min(top + float(shape.height) / deck.slide_height, 1.0) - y
            if x >= 1 or y >= 1 or w <= 0 or h <= 0:
                continue
            text = shape.text if shape.has_text_frame else ""
            if shape.has_table:
                text = "\n".join(
                    " | ".join(cell.text for cell in row.cells) for row in shape.table.rows
                )
            targets.append(
                Target(
                    id=f"shape_{shape.shape_id}",
                    text=text or shape.name,
                    kind=str(shape.shape_type),
                    is_text=bool(shape.has_text_frame and shape.text.strip()),
                    bounds=Bounds(x=x, y=y, width=w, height=h),
                ).model_dump(mode="json")
            )
        result[str(index)] = targets
    for item in text_lines or []:
        try:
            target = Target.model_validate(item["target"])
            slide_targets = result[str(item["slide"])]
            parent = next(t for t in slide_targets if t["id"] == target.parent_id)
            if (
                target.geometry_source != "powerpoint_text_line"
                or target.kind != "text_line"
                or not target.text
                or not target.id.startswith(f"{parent['id']}_line_")
                or any(t["id"] == target.id for t in slide_targets)
            ):
                raise ValueError("Invalid native text target")
            slide_targets.append(target.model_dump(mode="json"))
        except (KeyError, ValueError, TypeError, StopIteration):
            raise InputInvalid("PowerPoint returned invalid text-line geometry.") from None
    return result


def load_targets(root: Path, index: int) -> list[Target]:
    try:
        data = json.loads((root / "work/render/layout.json").read_text(encoding="utf-8"))
        return [Target.model_validate(item) for item in data[str(index)]]
    except (OSError, ValueError, KeyError, TypeError):
        raise InputInvalid(
            "Missing or invalid slide layout.", remedy="Run the render stage again."
        ) from None


def plan_path(root: Path, index: int) -> Path:
    return root / "work/script" / f"slide_{index:03}.plan.json"


def load_plan(root: Path, index: int) -> LecturePlan:
    try:
        return LecturePlan.model_validate_json(plan_path(root, index).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise InputInvalid(
            "Missing or invalid lecture plan.", remedy="Run the lecture_script stage again."
        ) from None


def compile_timeline(project: Project, root: Path, max_effects: int = 6) -> dict[str, Any]:
    slides = []
    for slide in project.slides:
        plan = load_plan(root, slide.index)
        targets = {target.id: target for target in load_targets(root, slide.index)}
        if len(plan.beats) != len(slide.segments):
            raise InputInvalid("Lecture beats and narration segments do not match.")
        beats = []
        effects = 0
        for beat, seg in zip(plan.beats, slide.segments, strict=True):
            if seg.text != beat.text or seg.start is None or seg.end is None:
                raise InputInvalid("Lecture plan does not match the synthesized narration.")
            effect = beat.effect
            bounds = None
            if effect != "none":
                target = targets.get(beat.target_id or "")
                if target is None:
                    raise InputInvalid("Lecture effect references an unknown slide target.")
                if effects >= max_effects:
                    effect = "none"
                else:
                    bounds = target.bounds.model_dump()
                    effects += 1
            beats.append(
                {
                    "segment_id": seg.id,
                    "text": seg.text,
                    "purpose": beat.purpose,
                    "start": seg.start,
                    "end": seg.end,
                    "effect": effect,
                    "target_id": beat.target_id,
                    "bounds": bounds,
                    "target_text": targets[beat.target_id or ""].text if bounds else None,
                    "target_kind": targets[beat.target_id or ""].kind if bounds else None,
                    "target_is_text": targets[beat.target_id or ""].is_text if bounds else False,
                    "geometry_source": targets[beat.target_id or ""].geometry_source
                    if bounds
                    else None,
                }
            )
        slides.append(
            {
                "index": slide.index,
                "title": slide.title,
                "start": slide.start,
                "end": slide.end,
                "objective": plan.objective,
                "takeaway": plan.takeaway,
                "warnings": plan.warnings,
                "beats": beats,
            }
        )
    return {
        "version": 2,
        "timing_basis": "sentence_audio_duration; not independently word-aligned",
        "human_listening_verified": False,
        "independent_asr_verified": False,
        "academic_content_verified": False,
        "slides": slides,
    }


def save_timeline(project: Project, root: Path, out: Path, max_effects: int) -> dict[str, Any]:
    timeline = compile_timeline(project, root, max_effects)
    write_json(out, timeline)
    return timeline
