"""Synthetic PNG renderer: extracts PPTX text, does not render slide graphics."""

import struct
import zlib
from pathlib import Path

from pptx import Presentation

from slidevoice.core.errors import InputInvalid
from slidevoice.core.interfaces import SlideRenderer
from slidevoice.core.models import Project, Slide
from slidevoice.core.registry import register


def color_png(path: Path, width: int, height: int, color: bytes) -> None:
    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    pixels = zlib.compress((b"\x00" + color * width) * height)
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", pixels) + chunk(b"IEND", b"")
    )


@register("render", "fake_render")
class FakeRenderer(SlideRenderer):
    name = "fake_render"

    def render(self, project: Project, out_dir: Path, width: int, height: int) -> list[Slide]:
        source = self.workdir / project.source_slides
        if source.suffix.lower() != ".pptx":
            raise InputInvalid(
                "fake_render supports PPTX only.",
                remedy="Use the generated samples/demo.pptx for M0.",
            )
        try:
            deck = Presentation(str(source))
        except Exception:
            raise InputInvalid(
                "Cannot open the PPTX file.",
                remedy="Verify the file opens in a presentation editor.",
            ) from None
        if not project.meta.title:
            project.meta.title = deck.core_properties.title or source.stem
        if not project.meta.author:
            project.meta.author = deck.core_properties.author
        result = []
        out_dir.mkdir(parents=True, exist_ok=True)
        for index, item in enumerate(deck.slides, 1):
            image = out_dir / f"slide_{index:03}.png"
            color_png(image, width, height, bytes((25, 40 + index % 100, 70)))
            text = "\n".join(shape.text for shape in item.shapes if shape.has_text_frame)
            notes = ""
            if item.has_notes_slide and item.notes_slide.notes_text_frame is not None:
                notes = item.notes_slide.notes_text_frame.text.strip()
            result.append(
                Slide(
                    index=index,
                    title=item.shapes.title.text if item.shapes.title else "",
                    body_text=text,
                    notes_raw=notes,
                    hidden=item._element.get("show", "1") == "0",
                    image_path=image.relative_to(self.workdir),
                )
            )
        if not result:
            raise InputInvalid("The deck contains no slides.")
        return result
