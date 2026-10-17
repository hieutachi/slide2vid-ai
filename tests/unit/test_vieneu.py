import json
from pathlib import Path

import httpx
import pytest

from samples.vietnamese_qa import make_deck, write_player
from slidevoice.core.errors import ConsentMissing, InputInvalid, PluginUnavailable, StageFailed
from slidevoice.core.media import write_tone
from slidevoice.core.models import VoiceProfile
from slidevoice.core.worker_client import WorkerClient
from slidevoice.plugins.tts.vieneu import VieNeuTTS
from workers.vieneu.server import handler_for, output_path


def profile():
    return VoiceProfile(
        ref_audio=Path("must-never-be-read.wav"),
        ref_text="unused",
        language="vi",
        consent_confirmed=True,
    )


def test_vieneu_consent_reset_and_language(tmp_path):
    engine = VieNeuTTS({}, tmp_path)
    with pytest.raises(ConsentMissing):
        engine.synthesize("Xin chào", tmp_path / "out.wav")
    voice = profile()
    engine.load_voice(voice)
    voice.consent_confirmed = False
    with pytest.raises(ConsentMissing):
        engine.load_voice(voice)
    with pytest.raises(ConsentMissing):
        engine.synthesize("Xin chào", tmp_path / "out.wav")
    voice.consent_confirmed = True
    voice.language = "en"
    with pytest.raises(PluginUnavailable):
        engine.load_voice(voice)


def test_vieneu_json_contract_has_no_reference_audio(tmp_path, monkeypatch):
    requests = []
    target = tmp_path / "speech.wav"

    def reply(request):
        if request.method == "GET":
            return httpx.Response(200, json={"status": "ready"})
        payload = json.loads(request.content)
        requests.append(payload)
        write_tone(target, 1)
        return httpx.Response(
            200,
            json={
                "out_path": "speech.wav",
                "sample_rate": 24000,
                "audible_signal": True,
            },
        )

    engine = VieNeuTTS({}, tmp_path)
    monkeypatch.setattr(
        engine,
        "_client",
        lambda: WorkerClient(
            "http://127.0.0.1:8770", tmp_path, transport=httpx.MockTransport(reply)
        ),
    )
    engine.healthcheck()
    engine.load_voice(profile())
    assert engine.synthesize("Xin chào Việt Nam.", target) == 1
    assert target.with_suffix(".txt").read_text(encoding="utf-8") == "Xin chào Việt Nam."
    assert requests[0]["voice"] == "Mai Anh"
    assert requests[0]["consent_confirmed"] is True
    assert not {"ref_audio", "ref_text", "base64"} & requests[0].keys()
    with pytest.raises(InputInvalid):
        engine.synthesize("Xin chào", tmp_path.parent / "outside.wav")
    for speed in (0, float("nan"), 3):
        with pytest.raises(InputInvalid):
            engine.synthesize("Xin chào", target, speed)


def test_vieneu_rejects_worker_contract_mismatch(tmp_path, monkeypatch):
    engine = VieNeuTTS({}, tmp_path)
    engine.load_voice(profile())
    transport = httpx.MockTransport(
        lambda _: httpx.Response(200, json={"out_path": "wrong.wav", "sample_rate": 48000})
    )
    monkeypatch.setattr(
        engine,
        "_client",
        lambda: WorkerClient("http://127.0.0.1:8770", tmp_path, transport=transport),
    )
    with pytest.raises(StageFailed, match="inconsistent"):
        engine.synthesize("Xin chào", tmp_path / "out.wav")


@pytest.mark.parametrize(
    "url", ["https://127.0.0.1:8770", "http://example.com", "http://192.168.1.2"]
)
def test_vieneu_rejects_remote_worker(tmp_path, url):
    with pytest.raises(InputInvalid):
        VieNeuTTS({"worker_url": url}, tmp_path).healthcheck()


def test_worker_path_and_consent_validation(tmp_path):
    payload = {
        "task": "tts_preset",
        "workdir": str(tmp_path / "job"),
        "out_path": "voice.wav",
        "consent_confirmed": True,
        "language": "vi",
        "text": "Xin chào",
    }
    assert output_path(payload, tmp_path) == tmp_path / "job/voice.wav"
    for change in (
        {"out_path": "../../escape.wav"},
        {"out_path": str(tmp_path / "absolute.wav")},
        {"workdir": str(tmp_path.parent)},
        {"out_path": "bad.exe"},
        {"consent_confirmed": False},
        {"consent_confirmed": "true"},
        {"text": ""},
        {"speed": float("inf")},
        {"task": "clone"},
    ):
        with pytest.raises(ValueError):
            output_path({**payload, **change}, tmp_path)


def test_worker_symlink_cannot_escape_allowed_root(tmp_path):
    root = tmp_path / "allowed"
    root.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    try:
        (root / "link").symlink_to(outside, target_is_directory=True)
    except OSError:
        pytest.skip("Creating symlinks requires Windows Developer Mode or privileges")
    with pytest.raises(ValueError):
        output_path(
            {
                "task": "tts_preset",
                "workdir": str(root),
                "out_path": "link/voice.wav",
                "consent_confirmed": True,
                "language": "vi",
                "text": "Xin chào",
            },
            root,
        )


def test_worker_http_routes_and_browser_rejection(tmp_path):
    import threading
    from http.server import HTTPServer

    class Engine:
        metadata = {"device": "test GPU"}
        voices = [("Test", "Test")]

        def run(self, body, root):
            output_path(body, root)
            return {"ok": True}

    server = HTTPServer(("127.0.0.1", 0), handler_for(Engine(), tmp_path))
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with httpx.Client(base_url=f"http://127.0.0.1:{server.server_port}") as client:
            assert client.get("/health").json()["status"] == "ready"
            assert len(client.get("/voices").json()["voices"]) == 1
            assert client.get("/absent").status_code == 404
            assert (
                client.post("/run", json={}, headers={"Origin": "http://evil.test"}).status_code
                == 403
            )
            assert client.post("/run", json={}).status_code == 400
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_vietnamese_fixture_is_not_english_and_player_is_honest(tmp_path):
    from pptx import Presentation

    deck_path = tmp_path / "slides.pptx"
    make_deck(deck_path)
    deck = Presentation(deck_path)
    assert len(deck.slides) == 3
    text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "TIẾNG VIỆT" in text
    assert "SYNTHETIC ENGLISH" not in text
    write_player(tmp_path, [{"voice": "<unsafe>", "native_path": "voice.wav"}])
    page = (tmp_path / "index.html").read_text(encoding="utf-8")
    assert "&lt;unsafe&gt;" in page
    assert "chưa có ASR độc lập" in page
    assert 'srclang="vi"' in page
