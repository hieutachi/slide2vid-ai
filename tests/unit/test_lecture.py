import json
from pathlib import Path

import httpx
import pytest
from pydantic import ValidationError
from typer.testing import CliRunner

from slidevoice.cli import app
from slidevoice.core.cache import write_json
from slidevoice.core.config import load_config
from slidevoice.core.errors import InputInvalid, PluginUnavailable, StageFailed
from slidevoice.core.pipeline import Pipeline
from slidevoice.core.presentation import (
    Bounds,
    LecturePlan,
    compile_timeline,
    extract_layout,
    load_targets,
    plan_path,
)
from slidevoice.plugins.llm import FakeLLM
from slidevoice.plugins.llm.openai_compat import OpenAICompat
from slidevoice.plugins.script.lecture import LectureScript, strict_schema


def response_for(user, effects=("none", "highlight", "circle", "zoom")):
    data = json.loads(user)
    target = data["targets"][0]["id"]
    quote = data["slide_text"] or data["notes"]
    return {
        "objective": "Hiểu nội dung chính.",
        "takeaway": "Nhớ ý chính.",
        "warnings": [],
        "beats": [
            {
                "text": f"Ý thứ {i}.",
                "purpose": "explain",
                "evidence": [
                    {"source": "slide" if data["slide_text"] else "notes", "quote": quote}
                ],
                "effect": effect,
                "target_id": None if effect == "none" else target,
            }
            for i, effect in enumerate(effects, 1)
        ],
    }


def fixture_llm(monkeypatch):
    monkeypatch.setattr(
        FakeLLM, "complete_json", lambda self, system, user, schema: response_for(user)
    )


def test_layout_and_validated_lecture(job, monkeypatch):
    fixture_llm(monkeypatch)
    project = Pipeline(job, {"script": {"plugin": "lecture_script"}}).run(until="script")
    slide = project.slides[0]
    targets = load_targets(job, slide.index)
    assert targets and targets[0].id.startswith("shape_")
    assert all(t.bounds.x + t.bounds.width <= 1.000001 for t in targets)
    plan = LecturePlan.model_validate_json(plan_path(job, slide.index).read_text(encoding="utf-8"))
    assert len(slide.segments) == len(plan.beats) == 4
    assert slide.script_source == "llm_from_notes"
    assert (job / "work/script/storyboard.md").is_file()
    again = Pipeline(job)
    again.run(until="script")
    assert again.events["script"] == "cache_hit"
    # A damaged sidecar must invalidate script, not silently survive cache reuse.
    plan_path(job, slide.index).write_text("{}", encoding="utf-8")
    rerun = Pipeline(job)
    rerun.run(until="script")
    assert rerun.events["script"] == "executed"


@pytest.mark.parametrize("problem", ["target", "quote", "extra", "budget", "bounds"])
def test_rejects_invalid_plans(job, monkeypatch, problem):
    project = Pipeline(job).run(until="render")
    slide = project.slides[0]

    def reply(self, system, user, schema):
        data = response_for(user)
        if problem == "target":
            data["beats"][1]["target_id"] = "nonexistent"
        elif problem == "quote":
            data["beats"][0]["evidence"][0]["quote"] = "Fabricated evidence."
        elif problem == "extra":
            data["unsupported"] = True
        elif problem == "bounds":
            data["beats"][0]["bounds"] = {"x": -1}
        return data

    monkeypatch.setattr(FakeLLM, "complete_json", reply)
    writer = LectureScript({"max_words_per_slide": 1 if problem == "budget" else 240}, job)
    with pytest.raises(StageFailed, match="validation failed"):
        writer.write(slide, "", FakeLLM({}, job))
    assert not plan_path(job, slide.index).exists()


def test_empty_slide_and_no_llm(job):
    from slidevoice.core.models import Slide

    writer = LectureScript({}, job)
    with pytest.raises(InputInvalid, match="requires an LLM"):
        writer.write(Slide(index=1), "", None)
    with pytest.raises(InputInvalid, match="no readable"):
        writer.write(Slide(index=1), "", FakeLLM({}, job))
    with pytest.raises(ValidationError):
        Bounds(x=0.9, y=0, width=0.2, height=0.2)


def test_timeline_matches_actual_audio_and_limits_effects(job, monkeypatch):
    fixture_llm(monkeypatch)
    project = Pipeline(job, {"script": {"plugin": "lecture_script"}}).run(until="tts")
    timeline = compile_timeline(project, job, max_effects=1)
    assert timeline["independent_asr_verified"] is False
    for slide, direction in zip(project.slides, timeline["slides"], strict=True):
        assert [b["effect"] for b in direction["beats"]] == ["none", "highlight", "none", "none"]
        for segment, beat in zip(slide.segments, direction["beats"], strict=True):
            assert beat["start"] == segment.start
            assert beat["end"] == segment.end
    project.slides[0].segments[0].text = "Changed after planning."
    with pytest.raises(InputInvalid, match="does not match"):
        compile_timeline(project, job)


def mock_http(monkeypatch, handler):
    original = httpx.Client
    monkeypatch.setenv("TEST_LECTURE_KEY", "sensitive-test-credential")
    monkeypatch.setattr(
        "slidevoice.plugins.llm.openai_compat.httpx.Client",
        lambda **kwargs: original(transport=httpx.MockTransport(handler), **kwargs),
    )
    monkeypatch.setattr("slidevoice.plugins.llm.openai_compat.time.sleep", lambda _: None)


def provider(tmp_path, **extra):
    return OpenAICompat(
        {"api_key_env": "TEST_LECTURE_KEY", "model": "test-model", **extra}, tmp_path
    )


@pytest.mark.parametrize("mode", ["json_schema", "json_object"])
def test_provider_structured_http_and_retry(tmp_path, monkeypatch, mode):
    requests = []

    def reply(request):
        requests.append(request)
        if len(requests) == 1:
            return httpx.Response(429)
        return httpx.Response(
            200,
            json={"choices": [{"finish_reason": "stop", "message": {"content": '{"ok": true}'}}]},
        )

    mock_http(monkeypatch, reply)
    result = provider(tmp_path, response_format=mode).complete_json(
        "system", "user", strict_schema()
    )
    assert result == {"ok": True}
    assert len(requests) == 2
    payload = json.loads(requests[-1].content)
    assert payload["response_format"]["type"] == mode
    assert requests[-1].headers["Authorization"] == "Bearer sensitive-test-credential"
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("failure", ["auth", "json", "truncated", "refusal", "secret", "transport"])
def test_provider_sanitized_failures(tmp_path, monkeypatch, failure):
    def reply(request):
        if failure == "transport":
            raise httpx.ConnectError("sensitive-test-credential", request=request)
        content = "sensitive-test-credential" if failure == "secret" else "not-json"
        if failure == "auth":
            return httpx.Response(401, text="sensitive-test-credential")
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "length" if failure == "truncated" else "stop",
                        "message": {"content": None if failure == "refusal" else content},
                    }
                ]
            },
        )

    mock_http(monkeypatch, reply)
    with pytest.raises(StageFailed) as exc:
        provider(tmp_path, max_retries=0).complete_json("", "", {})
    assert "sensitive-test-credential" not in str(exc.value)


def test_provider_requires_credentials_and_safe_endpoint(tmp_path, monkeypatch):
    monkeypatch.delenv("TEST_LECTURE_KEY", raising=False)
    with pytest.raises(PluginUnavailable):
        provider(tmp_path).healthcheck()
    monkeypatch.setenv("TEST_LECTURE_KEY", "credential")
    for url in (
        "http://remote.example/v1",
        "https://user:pass@example.com",
        "https://example.com/#token",
    ):
        with pytest.raises(InputInvalid):
            provider(tmp_path, base_url=url).healthcheck()


def test_lecture_cli_review_needs_no_voice_or_worker(job, monkeypatch):
    fixture_llm(monkeypatch)
    config = load_config(Path(__file__).resolve().parents[2] / "configs/lecture_vietnamese.yaml")
    config["render"]["plugin"] = "fake_render"
    config["llm"]["plugin"] = "fake_llm"
    path = job.parent / "lecture.json"
    # JSON is also valid YAML.
    write_json(path, config)
    out = job.parent / "review"
    result = CliRunner().invoke(
        app,
        [
            "lecture",
            "--slides",
            str(job / "inputs/slides.pptx"),
            "--out",
            str(out),
            "--config",
            str(path),
            "--review-only",
        ],
    )
    assert result.exit_code == 0, result.output
    assert (out / "work/script/storyboard.md").is_file()
    assert not (out / "work/tts").exists()


def test_layout_clips_shapes_at_slide_edges(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(-Inches(1), Inches(1), Inches(2), Inches(1))
    shape.text = "Partially outside"
    path = tmp_path / "clipped.pptx"
    deck.save(path)
    target = extract_layout(path)["1"][0]
    assert target["bounds"]["x"] == 0
    assert target["bounds"]["width"] == pytest.approx(Inches(1) / deck.slide_width)


def test_preset_placeholder_cannot_be_used_for_cloning():
    from slidevoice.core.config import merge, validate

    with pytest.raises(InputInvalid):
        validate(
            merge(load_config(), {"voice": {"plugin": "preset_voice"}, "tts": {"plugin": "xtts"}})
        )
