import json

import pytest
from pptx import Presentation

from samples.lecture_long_content import LESSONS
from samples.lecture_long_demo import LongFixture, make_deck
from slidevoice.core.presentation import LecturePlan, extract_layout


def test_long_fixture_has_24_distinct_slides_and_grounded_plans(tmp_path):
    path = tmp_path / "lesson.pptx"
    make_deck(path)
    deck = Presentation(path)
    assert len(deck.slides) == len(LESSONS) == 24
    assert len({lesson[0] for lesson in LESSONS}) == 24
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9, abs=0.000001)
    layout = extract_layout(path)
    fixture = LongFixture({}, tmp_path)
    for index, slide in enumerate(deck.slides, 1):
        assert slide.shapes.title.text == LESSONS[index - 1][0]
        raw = fixture.complete_json(
            "",
            json.dumps(
                {
                    "slide_index": index,
                    "targets": layout[str(index)],
                }
            ),
            {},
        )
        plan = LecturePlan.model_validate(raw)
        assert len(plan.beats) == 2
        for beat in plan.beats:
            assert beat.text in slide.notes_slide.notes_text_frame.text
            assert beat.effect != "zoom"
        assert any(shape.name == "number" for shape in slide.shapes)
