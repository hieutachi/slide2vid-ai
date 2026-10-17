import shutil
import struct
import subprocess

import pytest
from pptx import Presentation
from pptx.util import Inches

from samples.make_samples import make_samples
from slidevoice.core.config import load_config
from slidevoice.core.pipeline import Pipeline, create_job


@pytest.mark.skipif(
    not (shutil.which("soffice") or shutil.which("libreoffice")) or not shutil.which("pdftoppm"),
    reason="Real LibreOffice/Poppler executables are not installed",
)
def test_real_libreoffice_hidden_slide_notes_and_letterbox(tmp_path):
    slides, voice = make_samples(tmp_path / "fixtures")
    deck = Presentation(slides)
    deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
    deck.slides[1]._element.set("show", "0")
    deck.save(slides)
    config = load_config()
    config["render"].update(plugin="libreoffice", width=320, height=180)
    job = create_job(slides, voice, tmp_path / "job", config)
    project = Pipeline(job).run(until="render")
    assert [s.index for s in project.slides] == [1, 3]
    assert project.slides[0].notes_raw.startswith("Xin chào")
    assert project.slides[1].title == "Kết thúc"
    for slide in project.slides:
        assert struct.unpack(">II", (job / slide.image_path).read_bytes()[16:24]) == (320, 180)
        # A 4:3 deck at 16:9 output must have side padding, not stretched text.
        pixel = subprocess.run(
            [
                "ffmpeg",
                "-v",
                "error",
                "-nostdin",
                "-i",
                str(job / slide.image_path),
                "-vf",
                "crop=2:2:0:88,scale=1:1",
                "-frames:v",
                "1",
                "-pix_fmt",
                "rgb24",
                "-f",
                "rawvideo",
                "-",
            ],
            check=True,
            capture_output=True,
            timeout=30,
        ).stdout
        assert len(pixel) == 3 and max(pixel) < 5
    pipeline = Pipeline(job)
    pipeline.run(until="render")
    assert pipeline.events == {"render": "cache_hit"}
