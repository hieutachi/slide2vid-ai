import json
import time
from pathlib import Path
from unittest.mock import patch

import pytest
from pptx import Presentation
from typer.testing import CliRunner

from slidevoice.cli import app
from slidevoice.core.cache import file_hash
from slidevoice.core.errors import ConsentMissing, StageFailed
from slidevoice.core.media import measure_audio, run_process, wav_duration
from slidevoice.core.pipeline import Pipeline, create_job
from slidevoice.plugins.compose import FakeComposer
from slidevoice.plugins.tts import FakeTTS


def test_complete_package_resume_restart_and_consent(job):
    start = time.perf_counter()
    pipeline = Pipeline(job)
    project = pipeline.run()
    elapsed = time.perf_counter() - start
    assert elapsed < 60
    required = [
        "video.mp4",
        "subtitles.srt",
        "subtitles.vtt",
        "subtitles.ass",
        "transcript.md",
        "chapters.txt",
        "manifest.json",
        "report.json",
        "report.html",
    ]
    assert all((job / name).is_file() for name in required)
    report = json.loads((job / "report.json").read_text(encoding="utf-8"))
    assert report["stages"] == pipeline.events
    assert report["stage_seconds"] == pipeline.durations
    assert report["stage_seconds"].keys() == pipeline.events.keys()
    assert all(seconds >= 0 for seconds in report["stage_seconds"].values())
    assert report["qa_status"] == "passed"
    assert report["audio"]["audible_signal"] is True
    assert -18 < report["lufs"] < -14
    probe = json.loads(
        run_process(
            [
                "ffprobe",
                "-v",
                "error",
                "-show_streams",
                "-show_chapters",
                "-show_format",
                "-of",
                "json",
                str(job / "video.mp4"),
            ]
        )
    )
    streams = {s["codec_type"]: s for s in probe["streams"]}
    assert streams["video"]["codec_name"] == "h264"
    assert streams["audio"]["codec_name"] == "aac"
    assert streams["audio"]["sample_rate"] == "48000"
    assert streams["audio"]["channels"] == 2
    assert streams["audio"]["disposition"]["default"] == 1
    assert streams["subtitle"]["codec_name"] == "mov_text"
    assert streams["subtitle"]["tags"]["language"] == "vie"
    assert len(probe["chapters"]) == len(project.slides) == 3
    tags = probe["format"]["tags"]
    assert {"title", "artist", "language", "creation_time", "encoder", "comment"} <= tags.keys()
    assert "AI" in tags["comment"]
    assert tags["encoding_tool"].startswith("SlideVoice")
    assert streams["video"]["tags"]["encoder"].startswith("SlideVoice")
    for chapter, slide in zip(probe["chapters"], project.slides, strict=True):
        assert float(chapter["start_time"]) == pytest.approx(slide.start, abs=0.001)
        for seg in slide.segments:
            assert seg.end - seg.start == pytest.approx(wav_duration(job / seg.audio_path))
            assert seg.start <= seg.words[0].start <= seg.words[-1].end <= seg.end + 1e-6
    original = file_hash(job / "video.mp4")
    resumed = Pipeline(job)
    resumed.run()
    assert set(resumed.events.values()) == {"cache_hit"}
    assert file_hash(job / "video.mp4") == original
    with patch.object(FakeTTS, "synthesize", side_effect=AssertionError("cache expected")) as synth:
        restarted = Pipeline(job)
        restarted.run(from_stage="tts")
        synth.assert_not_called()
    assert restarted.events["voice"] == "cache_hit"
    assert restarted.events["tts"] == "executed"
    for kwargs in ({}, {"from_stage": "tts"}, {"from_stage": "compose"}):
        with pytest.raises(ConsentMissing), patch.object(FakeTTS, "synthesize") as synth:
            Pipeline(job, {"project": {"voice_consent": False}}).run(**kwargs)
            synth.assert_not_called()
    # Returning authorization to true permits normal use again, without consent copied from old cache.
    Pipeline(job, {"project": {"voice_consent": True}}).run()


def test_notes_change_only_resynthesizes_one_sentence(job, monkeypatch):
    Pipeline(job).run(until="tts")
    config = Pipeline(job).config
    slides = job.parent / "samples/demo.pptx"
    voice = job.parent / "samples/voice.wav"
    source_hash = file_hash(voice)
    deck = Presentation(slides)
    notes = deck.slides[0].notes_slide.notes_text_frame
    notes.text = notes.text.replace("Xin chào các bạn.", "Xin chào mọi người.")
    deck.save(slides)
    seen = []
    original = FakeTTS.synthesize

    def spy(self, text, out_path, speed=1.0):
        seen.append(text)
        return original(self, text, out_path, speed)

    monkeypatch.setattr(FakeTTS, "synthesize", spy)
    create_job(slides, voice, job, config)
    Pipeline(job).run(until="tts")
    assert seen == ["Xin chào mọi người."]
    assert file_hash(voice) == source_hash


def test_resume_after_failed_composition(job, monkeypatch):
    original = FakeComposer.compose

    def fail(*_):
        raise StageFailed("Injected compose failure")

    monkeypatch.setattr(FakeComposer, "compose", fail)
    with pytest.raises(StageFailed):
        Pipeline(job).run()
    assert (
        json.loads((job / "manifest.json").read_text(encoding="utf-8"))["stage_status"]["compose"]
        == "failed"
    )
    monkeypatch.setattr(FakeComposer, "compose", original)
    pipeline = Pipeline(job)
    pipeline.run()
    assert pipeline.events["tts"] == "cache_hit"
    assert pipeline.events["compose"] == "executed"


def test_missing_output_repairs_and_stale_restart_fails(job):
    Pipeline(job).run()
    (job / "video.mp4").unlink()
    pipeline = Pipeline(job)
    pipeline.run()
    assert pipeline.events["compose"] == "executed"
    assert (job / "video.mp4").is_file()
    (job / "work/render/slide_001.png").unlink()
    with pytest.raises(StageFailed, match="checkpoint"):
        Pipeline(job).run(from_stage="tts")
    pipeline = Pipeline(job)
    pipeline.run()
    assert pipeline.events["render"] == "executed"


def test_cli_run_resume_stage_preview_plugins(job):
    runner = CliRunner()
    source = job.parent / "samples"
    target = job.parent / "cli-job"
    fake = Path(__file__).resolve().parents[2] / "configs/fake.yaml"
    commands = [
        ["plugins"],
        [
            "run",
            "--slides",
            str(source / "demo.pptx"),
            "--voice",
            str(source / "voice.wav"),
            "--config",
            str(fake),
            "--out",
            str(target),
        ],
        ["resume", str(target)],
        [
            "run",
            "--slides",
            str(source / "demo.pptx"),
            "--voice",
            str(source / "voice.wav"),
            "--config",
            str(fake),
            "--out",
            str(target),
            "--from-stage",
            "tts",
        ],
        [
            "stage",
            "render",
            "--slides",
            str(source / "demo.pptx"),
            "--config",
            str(fake),
            "--out",
            str(job.parent / "render-only"),
        ],
        ["stage", "subtitle", "--job", str(target)],
        ["preview", "--slide", "2", str(target), "--out", str(job.parent / "preview")],
    ]
    for args in commands:
        result = runner.invoke(app, args)
        assert result.exit_code == 0, f"{args}: {result.output}: {result.exception}"
    manifest = json.loads((job.parent / "preview/manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["slides"]) == 1
    assert manifest["slides"][0]["index"] == 2
    assert manifest["slides"][0]["start"] == 0


def test_silent_final_output_is_reported_and_rejected(job):
    with pytest.raises(StageFailed, match="audible"):
        Pipeline(job, {"tts": {"fail": True}, "qa": {"max_retries": 0}}).run()
    report = json.loads((job / "report.json").read_text(encoding="utf-8"))
    assert report["qa_status"] == "failed"
    assert report["audio"]["audible_signal"] is False
    assert report["lufs"] is None
    assert report["warnings"]
    assert len(report["silent_segments"]) == 4
    # Cached silent media must never turn a failed QA run into a green resume.
    with pytest.raises(StageFailed, match="audible"):
        Pipeline(job).run()


def test_downloaded_video_composer_and_background_invalidation(job):
    background = job / "inputs/background.mp4"
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-f",
            "lavfi",
            "-i",
            "testsrc2=s=320x180:r=15:d=2",
            "-c:v",
            "libx264",
            "-preset",
            "ultrafast",
            "-pix_fmt",
            "yuv420p",
            str(background),
        ]
    )
    pipeline = Pipeline(job, {"compose": {"plugin": "qa_video"}})
    pipeline.run()
    assert measure_audio(job / "video.mp4")["audible_signal"] is True
    resumed = Pipeline(job)
    resumed.run()
    assert resumed.events["compose"] == "cache_hit"
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-f",
            "lavfi",
            "-i",
            "color=c=red:s=320x180:r=15:d=2",
            "-c:v",
            "libx264",
            "-preset",
            "ultrafast",
            "-pix_fmt",
            "yuv420p",
            str(background),
        ]
    )
    changed = Pipeline(job)
    changed.run()
    assert changed.events["compose"] == "executed"


def test_one_silent_sentence_fails_even_when_final_audio_has_signal(job, monkeypatch):
    from slidevoice.core.media import write_tone

    original = FakeTTS.synthesize

    def partly_silent(self, text, out_path, speed=1.0):
        duration = original(self, text, out_path, speed)
        if text == "Xin chào các bạn.":
            write_tone(out_path, duration, frequency=0)
        return duration

    monkeypatch.setattr(FakeTTS, "synthesize", partly_silent)
    with pytest.raises(StageFailed, match="audible"):
        Pipeline(job).run()
    report = json.loads((job / "report.json").read_text(encoding="utf-8"))
    assert report["audio"]["audible_signal"] is True
    assert report["silent_segments"] == ["s001_01"]
    assert report["qa_status"] == "failed"
