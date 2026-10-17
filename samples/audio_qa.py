"""Reproducible downloaded-video QA without real slides or speaker cloning.

Run python -m samples.audio_qa --source-video ABSOLUTE_MP4 --out ABSOLUTE_DIRECTORY.
Download the source separately; unit tests never require the network.
"""

import argparse
import shutil
from pathlib import Path

from pptx import Presentation

from samples.make_samples import make_samples
from slidevoice.core.cache import file_hash, write_json
from slidevoice.core.config import load_config
from slidevoice.core.media import measure_audio, run_process
from slidevoice.core.pipeline import Pipeline, create_job


def generate(source: Path, root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    source = source.resolve()
    source_levels = measure_audio(source)
    if not source_levels["audible_signal"]:
        raise ValueError("Downloaded source has no audible signal")
    assets = root / "fixtures"
    slides, voice = make_samples(assets)
    deck = Presentation()
    for title, text in (
        (
            "Audible speech check",
            "This is SlideVoice audio quality assurance. You should hear a clear synthetic voice, not silence or a test beep.",
        ),
        (
            "Downloaded video check",
            "The background is a public Sintel trailer from the Blender Foundation. The original soundtrack is not used for voice training or cloning.",
        ),
        (
            "Output package check",
            "This output includes stereo audio, subtitles, and three chapters. Please check your player volume and compare the separate audio file if playback is silent.",
        ),
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = text
    deck.save(slides)
    config_path = Path(__file__).resolve().parents[1] / "configs/qa_speech.yaml"
    job = create_job(slides, voice, root / "narrated", load_config(config_path))
    shutil.copyfile(source, job / "inputs/background.mp4")
    pipeline = Pipeline(job)
    pipeline.run()
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-i",
            str(job / "video.mp4"),
            "-map",
            "0:a:0",
            "-c:a",
            "pcm_s16le",
            str(job / "audio_check.wav"),
        ]
    )
    # Independent control: preserve the downloaded trailer's original audio and
    # entire video, normalizing only for playback. This is NOT pipeline narration.
    control = root / "source_audio_control.mp4"
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-i",
            str(source),
            "-map",
            "0:v:0",
            "-map",
            "0:a:0",
            "-c:v",
            "copy",
            "-c:a",
            "aac",
            "-ar",
            "48000",
            "-ac",
            "2",
            "-af",
            "loudnorm=I=-16:TP=-3:LRA=11",
            "-disposition:a:0",
            "default",
            "-movflags",
            "+faststart",
            str(control),
        ]
    )
    attribution = "Sintel: (c) copyright Blender Foundation | durian.blender.org\nSource: https://media.w3.org/2010/05/sintel/trailer.mp4\nLicense: CC BY 3.0 https://durian.blender.org/sharing/\nQA adaptation: synthetic English system narration; not voice cloning.\nControl: full source trailer with normalized original soundtrack.\n"
    (root / "ATTRIBUTION.txt").write_text(attribution, encoding="utf-8")
    shutil.copyfile(root / "ATTRIBUTION.txt", job / "ATTRIBUTION.txt")
    write_json(
        root / "comparison.json",
        {
            "source_sha256": file_hash(source),
            "source_audio": source_levels,
            "control_audio": measure_audio(control),
            "narration_audio": measure_audio(job / "video.mp4"),
            "narration_sha256": file_hash(job / "video.mp4"),
            "human_listening_verified": False,
        },
    )
    (root / "listen.html").write_text(
        '<!doctype html><meta charset="utf-8"><title>SlideVoice audio QA</title>'
        "<h1>SlideVoice audible QA</h1><p>Press Play and check that the player is not muted. No autoplay.</p>"
        "<h2>1. Synthetic English narration over downloaded footage</h2>"
        '<video controls preload="metadata" width="960" src="narrated/video.mp4">'
        '<track kind="subtitles" src="narrated/subtitles.vtt" srclang="en" label="English" default>'
        "</video>"
        '<h2>2. Same narration as PCM WAV</h2><audio controls src="narrated/audio_check.wav"></audio>'
        "<h2>3. Downloaded trailer: original soundtrack control</h2>"
        '<video controls preload="metadata" width="960" src="source_audio_control.mp4"></video>'
        "<p>Footage: (c) Blender Foundation | durian.blender.org, CC BY 3.0. See ATTRIBUTION.txt.</p>",
        encoding="utf-8",
    )
    return job


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-video", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    generate(args.source_video, args.out.resolve())
