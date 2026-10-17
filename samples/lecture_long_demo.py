"""24-slide Full-HD scale trial, with explicit fixture provenance and real TTS options."""

from __future__ import annotations

import argparse
import html
import json
import logging
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from samples.lecture_long_content import LESSONS
from slidevoice.core.cache import stable_hash, write_json
from slidevoice.core.config import load_config
from slidevoice.core.interfaces import LLMProvider
from slidevoice.core.media import write_tone
from slidevoice.core.pipeline import Pipeline, create_job
from slidevoice.core.registry import register

BG, WHITE, MUTED, TEAL = (17, 30, 51), (243, 247, 255), (158, 180, 204), (91, 207, 205)


def textbox(
    slide: Any,
    name: str,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: int,
    color: tuple[int, int, int] = WHITE,
) -> None:
    if name == "heading" and slide.shapes.title is not None:
        shape = slide.shapes.title
        shape.left, shape.top, shape.width, shape.height = (
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
        )
    else:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.text_frame.word_wrap = True
    shape.text_frame.text = text
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor(*color)
        paragraph.space_after = Pt(16)


def make_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = "Đọc hiểu dữ liệu — Bài thử 24 slide"
    for index, (title, lines, narration, values) in enumerate(LESSONS, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*BG)
        section = ("ĐẶT CÂU HỎI", "MÔ TẢ TRUNG TÂM", "ĐỌC PHÂN BỐ", "KẾT LUẬN CẨN TRỌNG")[
            (index - 1) // 6
        ]
        textbox(slide, "section", f"SLIDEVOICE / {section}", 0.7, 0.35, 11.8, 0.45, 15, TEAL)
        textbox(slide, "heading", title, 0.7, 1.1, 12, 1.05, 33)
        body_width = 7.6 if values else 11.8
        for i, line in enumerate(lines):
            textbox(slide, f"point_{i}", line, 0.7, 2.65 + i * 0.95, body_width, 0.85, 25)
        if values:
            chart_x, chart_y, chart_w, chart_h = 9.0, 2.8, 3.55, 2.9
            maximum = max(values)
            textbox(
                slide,
                "chart_caption",
                "MINH HỌA • ĐƠN VỊ GIẢ ĐỊNH",
                chart_x,
                2.2,
                chart_w,
                0.5,
                11,
                MUTED,
            )
            for i, value in enumerate(values):
                height = max(0.015, chart_h * value / maximum)
                x = chart_x + i * chart_w / len(values)
                bar = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    Inches(x),
                    Inches(chart_y + chart_h - height),
                    Inches(0.42),
                    Inches(height),
                )
                bar.fill.solid()
                bar.fill.fore_color.rgb = RGBColor(*TEAL)
                bar.line.fill.background()
                textbox(
                    slide,
                    f"value_{i}",
                    str(value),
                    x - 0.03,
                    chart_y + chart_h - height - 0.4,
                    0.72,
                    0.4,
                    13,
                    WHITE,
                )
                if index in {18, 19}:
                    textbox(slide, f"group_{i}", chr(65 + i), x, 5.74, 0.65, 0.4, 12, MUTED)
            textbox(
                slide,
                "chart_baseline",
                "Gốc = 0 • Nhóm A–E" if index in {18, 19} else "Gốc = 0 • Cột theo thứ tự dữ liệu",  # noqa: RUF001
                chart_x,
                5.9,
                chart_w,
                0.5,
                11,
                MUTED,
            )
        textbox(
            slide,
            "footer",
            "VÍ DỤ GIẢ ĐỊNH • GIỌNG AI • KỊCH BẢN SOẠN SẴN",
            0.7,
            6.65,
            11.8,
            0.4,
            12,
            MUTED,
        )
        textbox(slide, "number", f"{index:02} / {len(LESSONS)}", 11.6, 0.35, 1.2, 0.45, 14, TEAL)
        progress = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.7),
            Inches(7.2),
            Inches(11.9 * index / len(LESSONS)),
            Inches(0.035),
        )
        progress.fill.solid()
        progress.fill.fore_color.rgb = RGBColor(*TEAL)
        progress.line.fill.background()
        slide.notes_slide.notes_text_frame.text = "\n".join(narration)
    deck.save(path)


@register("llm", "lecture_long_fixture")
class LongFixture(LLMProvider):
    name = "lecture_long_fixture"
    version = "0.1.0"

    def complete_json(self, system: str, user: str, schema: dict[str, Any]) -> dict[str, Any]:
        data = json.loads(user)
        title, lines, narration, _ = LESSONS[data["slide_index"] - 1]
        targets = data["targets"]
        chosen = []
        for line in lines:
            matches = [target for target in targets if target["text"] == line]
            measured = [
                target for target in matches if target["geometry_source"] == "powerpoint_text_line"
            ]
            if not matches:
                raise ValueError(f"Missing fixture target: {line}")
            chosen.append((measured or matches)[0]["id"])
        return {
            "objective": title,
            "takeaway": " ".join(lines),
            "warnings": ["Bài thử 24 slide: kịch bản soạn sẵn, không gọi LLM API."],
            "beats": [
                {
                    "text": text,
                    "purpose": "orient" if i == 0 else "explain",
                    "evidence": [{"source": "notes", "quote": text}],
                    "effect": "none"
                    if i == 0
                    else ("circle" if data["slide_index"] % 3 == 0 else "highlight"),
                    "target_id": None if i == 0 else chosen[-1],
                }
                for i, text in enumerate(narration)
            ],
        }


def generate(root: Path, real: bool = False) -> None:
    root.mkdir(parents=True, exist_ok=True)
    deck, reference = root / "lesson.pptx", root / "preset_placeholder.wav"
    # Reusing existing input bytes preserves render-cache fingerprints on resume.
    fingerprint = stable_hash({"deck_version": "0.1.2", "lessons": LESSONS})
    stamp = root / "deck_fingerprint.txt"
    if not deck.exists() or not stamp.exists() or stamp.read_text() != fingerprint:
        make_deck(deck)
        stamp.write_text(fingerprint)
    if not reference.exists():
        write_tone(reference, 0.1, frequency=0)
    config = load_config(Path(__file__).resolve().parents[1] / "configs/lecture_vietnamese.yaml")
    config["llm"]["plugin"] = "lecture_long_fixture"
    config["project"]["title"] = "Đọc hiểu dữ liệu — 24 slide thử nghiệm"
    config["tts"]["plugin"] = "vieneu" if real else "fake_tts"
    config["render"]["plugin"] = "powerpoint" if real else "fake_render"
    config["render"].update(width=1920 if real else 640, height=1080 if real else 360)
    config["compose"].update(preset="fast", crf=18, fps=30 if real else 15)
    job = create_job(deck, reference, root / "lesson", config)
    Pipeline(job).run()
    write_json(
        root / "demo_evidence.json",
        {
            "llm": "hand-authored fixture; no real LLM API call",
            "slides": len(LESSONS),
            "speech": config["tts"]["plugin"],
            "render": config["render"]["plugin"],
            "human_listening_verified": False,
            "independent_asr_verified": False,
            "academic_content_verified": False,
        },
    )
    template = Path(__file__).with_name("lecture_player.html").read_text(encoding="utf-8")
    notice = "Giọng VieNeu dựng sẵn" if real else "Âm beep kiểm thử, KHÔNG phải giọng nói"
    (root / "index.html").write_text(
        template.replace("{{NOTICE}}", html.escape(notice)), encoding="utf-8"
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument(
        "--real", action="store_true", help="PowerPoint export + local VieNeu speech"
    )
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
    generate(args.out.resolve(), args.real)
