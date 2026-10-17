"""Reproducible Vietnamese preset listening samples and actual PowerPoint video."""

import argparse
import html
import json
import shutil
import urllib.request
from pathlib import Path

from samples import slides_qa
from slidevoice.core.cache import file_hash, write_json
from slidevoice.core.config import load_config
from slidevoice.core.media import measure_audio, run_process, wav_duration, write_tone
from slidevoice.core.models import VoiceProfile
from slidevoice.core.pipeline import Pipeline, create_job
from slidevoice.plugins.tts.vieneu import VieNeuTTS

ROOT = Path(__file__).resolve().parents[1]
SLIDES = (
    (
        "SLIDE 01 / GIỌNG ĐỌC TIẾNG VIỆT",
        "Một giọng đọc rõ ràng, tự nhiên.",
        "Mô hình chạy ngay trên máy của bạn.\nGiọng dựng sẵn, không nhân bản người dùng.",
        "Xin chào bạn, đây là bản thử giọng đọc tiếng Việt của SlideVoice. "
        "Mình sẽ đọc chậm rãi, ngắt nghỉ rõ ràng, để bạn dễ theo dõi nội dung trên từng trang trình chiếu.",
        (18, 46, 85),
    ),
    (
        "SLIDE 02 / KIỂM TRA PHÁT ÂM",
        "Dấu tiếng Việt và cách ngắt câu.",
        "Buổi sáng trong lành, nắng nhẹ qua ô cửa.\nLời đọc nối tiếp đúng hình ảnh của slide.",
        "Buổi sáng, ánh nắng nhẹ xuyên qua ô cửa, những hàng cây khẽ đung đưa trong gió. "
        "Dù ở Hà Nội, Đà Nẵng hay thành phố Hồ Chí Minh, một lời chào ấm áp luôn khiến chúng ta thấy gần nhau hơn.",
        (15, 70, 62),
    ),
    (
        "SLIDE 03 / MỜI BẠN NGHE THỬ",
        "Bạn thích chất giọng này chứ?",
        "So sánh bản thu gốc với âm thanh trong video.\nChỉ tinh chỉnh tiếp khi có tiêu chí và dữ liệu phù hợp.",
        "Bản thử này sử dụng mô hình đã được huấn luyện sẵn và xử lý bằng cạc đồ họa trên máy. "
        "Bạn hãy nghe thử chất giọng, độ rõ chữ và nhịp đọc, rồi cho mình biết phần nào cần điều chỉnh nhé.",
        (61, 35, 88),
    ),
)
COMPARISON_TEXT = (
    "Xin chào bạn, mình là giọng đọc tổng hợp tiếng Việt. "
    "Một bài thuyết trình hay không chỉ có hình ảnh đẹp, mà còn cần lời dẫn rõ ràng, "
    "ấm áp và đúng nhịp. Bạn thấy chất giọng này có phù hợp để đọc bài giảng không?"
)


def make_deck(path: Path) -> None:
    # Reuse the existing fixture's PowerPoint layout, replacing all English copy.
    original = slides_qa.SLIDES
    try:
        slides_qa.SLIDES = SLIDES
        slides_qa.make_deck(path)
    finally:
        slides_qa.SLIDES = original
    from pptx import Presentation

    deck = Presentation(path)
    deck.core_properties.title = "SlideVoice — giọng đọc tiếng Việt"
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if "SYNTHETIC ENGLISH QA" in run.text:
                            run.text = "SLIDEVOICE  |  GIỌNG TỔNG HỢP TIẾNG VIỆT"
    deck.save(path)


def write_player(root: Path, samples: list[dict[str, str]]) -> None:
    cards = "\n".join(
        f'<h3>{html.escape(sample["voice"])}</h3><audio controls preload="metadata" '
        f'src="{html.escape(sample["native_path"], quote=True)}"></audio>'
        for sample in samples
    )
    (root / "index.html").write_text(
        """<!doctype html><html lang="vi"><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>SlideVoice — nghe thử giọng Việt</title>
<style>body{max-width:1000px;margin:30px auto;padding:18px;font:18px system-ui;
background:#101c30;color:#eef4ff}video,audio{width:100%}a{color:#8edaff}
pre{white-space:pre-wrap}button{padding:12px;font:inherit}</style>
<h1>Giọng Việt + slide PowerPoint thật</h1>
<p>Giọng chính: <strong>Mai Anh</strong> — preset VieNeu-TTS v3 Turbo.
Đây là lời đọc do AI tạo, không phải bản thu của người dùng hay giọng được clone.</p>
<p>Chọn đúng tai nghe và bắt đầu ở mức âm lượng thấp. Trang không tự phát,
không thay đổi âm lượng hay thiết bị Windows.</p>
<video id="v" controls playsinline preload="metadata" src="slides_compatible.mp4">
<track kind="subtitles" src="narrated/subtitles.vtt" srclang="vi" label="Tiếng Việt">
</video><button id="play">Phát video</button><pre id="status">Đang nạp…</pre>
<h2>Lời đọc tách khỏi video</h2>
<audio controls preload="metadata" src="narrated/audio_check.wav"></audio>
<h2>So sánh chất giọng — WAV gốc 48 kHz, cùng một đoạn văn</h2>
"""
        + cards
        + """<p><a download href="slides_compatible.mp4">Tải video MP4</a> ·
<a download href="slides_fallback.webm">WebM dự phòng</a> ·
<a href="fixtures/slides.pptx">Slide PowerPoint</a> ·
<a href="qa_checks.json">Kết quả kỹ thuật</a></p>
<p>Phụ đề/căn chỉnh hiện là fixture dựa trên kịch bản, chưa có ASR độc lập.
Kiểm tra tín hiệu không chứng minh phát âm đúng hoặc giọng hay; cần bạn nghe đánh giá.</p>
<script>const v=document.querySelector('#v'),s=document.querySelector('#status');
function report(e){s.textContent=JSON.stringify({event:e.type,time:v.currentTime,
duration:v.duration,readyState:v.readyState,muted:v.muted,volume:v.volume,
error:v.error?{code:v.error.code,message:v.error.message}:null},null,2)}
['loadedmetadata','canplay','playing','pause','waiting','stalled','error','timeupdate']
.forEach(e=>v.addEventListener(e,report));
document.querySelector('#play').onclick=async()=>{try{v.muted=false;await v.play()}
catch(e){s.textContent=e.name+': '+e.message}};</script></html>""",
        encoding="utf-8",
    )


def generate(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    fixtures = root / "fixtures"
    fixtures.mkdir(exist_ok=True)
    slides = fixtures / "slides.pptx"
    placeholder = fixtures / "unused_reference.wav"
    make_deck(slides)
    # Pipeline currently requires a reference-file field, but the preset worker
    # never receives or reads it. This is not voice cloning or a training sample.
    write_tone(placeholder, 1)
    config = load_config(ROOT / "configs/qa_vietnamese.yaml")
    with urllib.request.urlopen(config["tts"]["worker_url"] + "/health", timeout=10) as response:
        health = json.load(response)
    with urllib.request.urlopen(config["tts"]["worker_url"] + "/voices", timeout=10) as response:
        voices = json.load(response)
    write_json(root / "worker_health.json", health)
    write_json(root / "available_voices.json", voices)
    samples = []
    for index, voice in enumerate(("Mai Anh", "Hải Đăng", "Thùy Dung"), 1):
        engine = VieNeuTTS({**config["tts"], "preset_voice": voice}, root)
        engine.load_voice(
            VoiceProfile(
                ref_audio=placeholder, ref_text="unused", language="vi", consent_confirmed=True
            )
        )
        path = root / "voice_samples" / f"voice_{index:02}.wav"
        engine.synthesize(COMPARISON_TEXT, path)
        samples.append(
            {
                "voice": voice,
                "native_path": path.with_suffix(".native48k.wav").relative_to(root).as_posix(),
            }
        )
    job = create_job(slides, placeholder, root / "narrated", config)
    project = Pipeline(job).run()
    compatible = root / "slides_compatible.mp4"
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-i",
            str(job / "video.mp4"),
            "-map",
            "0:v:0",
            "-map",
            "0:a:0",
            "-map_metadata",
            "-1",
            "-map_chapters",
            "-1",
            "-c:v",
            "libx264",
            "-profile:v",
            "baseline",
            "-level:v",
            "3.1",
            "-preset",
            "fast",
            "-crf",
            "20",
            "-pix_fmt",
            "yuv420p",
            "-c:a",
            "copy",
            "-movflags",
            "+faststart",
            str(compatible),
        ],
        timeout=300,
    )
    fallback = root / "slides_fallback.webm"
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-i",
            str(compatible),
            "-c:v",
            "libvpx-vp9",
            "-deadline",
            "realtime",
            "-cpu-used",
            "8",
            "-crf",
            "35",
            "-b:v",
            "0",
            "-c:a",
            "libopus",
            "-b:a",
            "128k",
            str(fallback),
        ],
        timeout=300,
    )
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-i",
            str(compatible),
            "-map",
            "0:a:0",
            "-c:a",
            "pcm_s16le",
            str(job / "audio_check.wav"),
        ]
    )
    for video in (compatible, fallback):
        run_process(
            [
                "ffmpeg",
                "-v",
                "error",
                "-xerror",
                "-nostdin",
                "-i",
                str(video),
                "-map",
                "0:v:0",
                "-map",
                "0:a:0",
                "-f",
                "null",
                "-",
            ],
            timeout=300,
        )
    audio = measure_audio(compatible)
    if not audio["audible_signal"]:
        raise RuntimeError("Final Vietnamese video has no audible signal.")
    metrics = [
        json.loads(p.read_text(encoding="utf-8")) for p in sorted(root.rglob("*.metrics.json"))
    ]
    write_json(
        root / "qa_checks.json",
        {
            "model": health,
            "main_voice": "Mai Anh",
            "voice_samples": samples,
            "slides": len(project.slides),
            "audio": audio,
            "duration_seconds": wav_duration(job / "audio_check.wav"),
            "generation_metrics": metrics,
            "full_decode_passed": [compatible.name, fallback.name],
            "human_listening_verified": False,
            "independent_asr_verified": False,
            "alignment": "synthetic-sidecar fixture; not ASR/forced alignment",
            "training_performed": False,
            "warnings": project.warnings,
        },
    )
    write_player(root, samples)
    shutil.copyfile(job / "work/render/slide_001.png", root / "slide_preview.png")
    write_json(
        root / "SHA256SUMS.json",
        {
            p.relative_to(root).as_posix(): file_hash(p)
            for p in sorted(root.rglob("*"))
            if p.is_file() and p.name != "SHA256SUMS.json"
        },
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    generate(parser.parse_args().out.resolve())
