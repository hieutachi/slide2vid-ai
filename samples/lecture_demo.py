"""Reproducible teaching-video demo with a labelled, hand-authored LLM fixture.

Use the lecture CLI for real LLM API input. This fixture tests local speech and
presentation without sending data or spending API credits.
"""

from __future__ import annotations

import argparse
import html
import json
import shutil
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from slidevoice.core.cache import write_json
from slidevoice.core.config import load_config
from slidevoice.core.interfaces import LLMProvider
from slidevoice.core.media import write_tone
from slidevoice.core.pipeline import Pipeline, create_job
from slidevoice.core.registry import register

LESSONS = (
    {
        "title": "Trung bình và trung vị",
        "body": "Hai cách mô tả giá trị trung tâm.\nDữ liệu minh họa: 2, 3, 4, 5, 6.",
        "beats": [
            (
                "Trước khi tính toán, chúng ta cần hiểu câu hỏi: một giá trị nào có thể đại diện cho cả nhóm dữ liệu?",
                "orient",
                "none",
            ),
            (
                "Trung bình và trung vị đều mô tả giá trị trung tâm, nhưng cách tính khác nhau. Vì thế, chúng có thể trả lời câu hỏi này theo những cách khác nhau.",
                "explain",
                "highlight",
            ),
            (
                "Ví dụ minh họa, với dãy hai, ba, bốn, năm, sáu, cả trung bình và trung vị đều bằng bốn. Chúng ta sẽ xem điều gì xảy ra khi thay đổi một giá trị.",
                "example",
                "circle",
            ),
        ],
    },
    {
        "title": "Khi xuất hiện giá trị ngoại lệ",
        "body": "Dữ liệu mới: 2, 3, 4, 5, 100.\nTrung bình = 22,8. Trung vị = 4.",
        "beats": [
            (
                "Bây giờ, thay số sáu bằng một trăm, trong khi bốn số còn lại được giữ nguyên.",
                "orient",
                "highlight",
            ),
            (
                "Tổng năm số mới là một trăm mười bốn. Chia cho năm, ta được trung bình bằng hai mươi hai phẩy tám. Giá trị một trăm đã kéo trung bình tăng lên đáng kể.",
                "explain",
                "zoom",
            ),
            (
                "Nhưng số đứng giữa dãy đã sắp xếp vẫn là bốn, nên trung vị không đổi. Đây là điểm khác biệt cần chú ý trong ví dụ này.",
                "explain",
                "circle",
            ),
        ],
    },
    {
        "title": "Chọn thước đo theo câu hỏi",
        "body": "Trung bình: dùng thông tin của mọi giá trị.\nTrung vị: vị trí giữa của dữ liệu đã sắp xếp.\nKhông có thước đo luôn tốt nhất.",
        "beats": [
            (
                "Trung bình sử dụng thông tin của mọi giá trị. Trung vị dựa vào vị trí giữa sau khi sắp xếp dữ liệu.",
                "recap",
                "highlight",
            ),
            (
                "Không nên kết luận rằng trung vị luôn tốt hơn trung bình. Hãy chọn thước đo dựa trên câu hỏi và đặc điểm của dữ liệu.",
                "explain",
                "circle",
            ),
            (
                "Để tự kiểm tra, bạn hãy đổi số một trăm thành một nghìn. Đại lượng nào thay đổi, và đại lượng nào vẫn giữ nguyên? Hãy giải thích bằng cách tính của từng thước đo.",
                "transition",
                "none",
            ),
        ],
    },
)


def make_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = "Hiểu trung bình và trung vị"
    for index, lesson in enumerate(LESSONS, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(17, 30, 51)
        for name, top, height, text, size, color in (
            ("label", 0.45, 0.45, f"SLIDEVOICE / THỐNG KÊ CƠ BẢN / {index:02}", 16, (91, 207, 205)),
            ("heading", 1.25, 1.1, lesson["title"], 36, (243, 247, 255)),
            ("content", 3.0, 2.3, lesson["body"], 28, (222, 232, 246)),
            (
                "footer",
                6.65,
                0.4,
                "VÍ DỤ MINH HỌA • GIỌNG AI • KỊCH BẢN DEMO SOẠN SẴN",
                13,
                (144, 167, 192),
            ),
        ):
            shape = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(height))
            shape.name = name
            shape.text_frame.word_wrap = True
            shape.text_frame.text = str(text)
            for paragraph in shape.text_frame.paragraphs:
                if name == "content":
                    paragraph.space_after = Pt(16)
                paragraph.font.name = "Arial"
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor(*color)
        slide.notes_slide.notes_text_frame.text = "\n".join(beat[0] for beat in lesson["beats"])
    deck.save(path)


@register("llm", "lecture_demo_fixture")
class DemoLLM(LLMProvider):
    name = "lecture_demo_fixture"
    version = "0.2.0"

    def complete_json(self, system: str, user: str, schema: dict[str, Any]) -> dict[str, Any]:
        data = json.loads(user)
        lesson = LESSONS[data["slide_index"] - 1]
        content = next(target for target in data["targets"] if target["text"] == lesson["body"])
        lines = [target for target in data["targets"] if target.get("parent_id") == content["id"]]
        # Fixture decisions are explicit, not claims of semantic AI targeting.
        selections = {1: (0, 0, 1), 2: (0, 1, 1), 3: (0, 2, 2)}[data["slide_index"]]
        return {
            "objective": lesson["title"],
            "takeaway": lesson["body"],
            "warnings": ["Kịch bản fixture soạn sẵn, không phải kết quả gọi LLM API."],
            "beats": [
                {
                    "text": text,
                    "purpose": purpose,
                    "evidence": [{"source": "notes", "quote": text}],
                    "effect": effect,
                    "target_id": None
                    if effect == "none"
                    else (
                        lines[min(selections[i], len(lines) - 1)]["id"] if lines else content["id"]
                    ),
                }
                for i, (text, purpose, effect) in enumerate(lesson["beats"])
            ],
        }


def generate(
    root: Path,
    speech: str = "fake_tts",
    render: str = "fake_render",
    reuse_audio: Path | None = None,
) -> None:
    root.mkdir(parents=True, exist_ok=True)
    slides, reference = root / "lesson.pptx", root / "preset_placeholder.wav"
    make_deck(slides)
    write_tone(reference, 0.1, frequency=0)
    cfg = load_config(Path(__file__).resolve().parents[1] / "configs/lecture_vietnamese.yaml")
    cfg["llm"]["plugin"] = "lecture_demo_fixture"
    cfg["tts"]["plugin"] = speech
    cfg["render"]["plugin"] = render
    cfg["render"].update(width=1280, height=720)
    job = create_job(slides, reference, root / "lesson", cfg)
    if reuse_audio is not None:
        source_cache = reuse_audio.resolve() / "work/tts/sentences"
        if not source_cache.is_dir():
            raise ValueError("Reuse source must be an existing job with a sentence cache.")
        # SegmentCache checks content hashes and full TTS/voice fingerprints before use.
        shutil.copytree(source_cache, job / "work/tts/sentences", dirs_exist_ok=True)
    project = Pipeline(job).run()
    notice = (
        "Giọng VieNeu dựng sẵn" if speech == "vieneu" else "Âm kiểm thử beep, KHÔNG phải giọng nói"
    )
    write_json(
        root / "demo_evidence.json",
        {
            "llm": "hand-authored fixture; no real LLM API call",
            "speech": speech,
            "render": render,
            "human_listening_verified": False,
            "independent_asr_verified": False,
            "academic_content_verified": False,
            "slides": len(project.slides),
        },
    )
    template = Path(__file__).with_name("lecture_player.html").read_text(encoding="utf-8")
    (root / "index.html").write_text(
        template.replace("{{NOTICE}}", html.escape(notice)), encoding="utf-8"
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--speech", choices=["fake_tts", "vieneu"], default="fake_tts")
    parser.add_argument("--render", choices=["fake_render", "powerpoint"], default="fake_render")
    parser.add_argument(
        "--reuse-audio",
        type=Path,
        help="Existing job; reuse only validated matching sentence audio",
    )
    args = parser.parse_args()
    generate(args.out.resolve(), args.speech, args.render, args.reuse_audio)
