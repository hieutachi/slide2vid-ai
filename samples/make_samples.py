"""Generate a three-slide Vietnamese deck and consent-free synthetic tone sample."""

from pathlib import Path

from pptx import Presentation

from slidevoice.core.media import write_tone


def make_samples(root: Path) -> tuple[Path, Path]:
    root.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.core_properties.title = "SlideVoice mẫu"
    deck.core_properties.author = "SlideVoice"
    for title, notes in (
        ("Giới thiệu", "Xin chào các bạn. Đây là bản thử nghiệm."),
        ("Quy trình", "Chúng ta kiểm tra từng bước."),
        ("Kết thúc", "Cảm ơn các bạn đã theo dõi."),
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "Ví dụ tiếng Việt có dấu — không có giọng người thật."
        slide.notes_slide.notes_text_frame.text = notes
    slides, voice = root / "demo.pptx", root / "voice.wav"
    deck.save(slides)
    write_tone(voice, 10, frequency=220)
    return slides, voice


if __name__ == "__main__":
    make_samples(Path(__file__).resolve().parent)
