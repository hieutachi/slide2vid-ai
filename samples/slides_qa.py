"""Real PowerPoint slides, system narration, browser-compatible playback package."""

import argparse
import json
import shutil
import wave
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from samples.playback import write_player
from slidevoice.core.cache import file_hash, write_json
from slidevoice.core.config import load_config
from slidevoice.core.media import measure_audio, run_process, write_tone
from slidevoice.core.models import VoiceProfile
from slidevoice.core.pipeline import Pipeline, create_job
from slidevoice.plugins.tts import WindowsSAPI

SLIDES = (
    (
        "SLIDE 01 / AUDIO CHECK",
        "Real slide. Spoken narration.",
        "PowerPoint graphics, not a blank background.\nEnglish system voice, not a test beep.",
        "This is slide one. You should see a blue slide and hear spoken narration. This is a system voice, not a test beep.",
        (18, 46, 85),
    ),
    (
        "SLIDE 02 / TIMELINE CHECK",
        "Picture and voice move together.",
        "Each slide has its own narration.\nSlide changes follow the audio timeline.",
        "This is slide two. The slide background is now green. Each slide has its own narration, and the pictures follow the audio timeline.",
        (15, 70, 62),
    ),
    (
        "SLIDE 03 / HEADPHONE CHECK",
        "Compare video, WAV, left and right.",
        "Use the playback page to compare audio.\nHuman listening confirmation is still required.",
        "This is slide three. The background is now purple. Please compare the video with the separate wave audio, then try the left and right headphone check.",
        (61, 35, 88),
    ),
)


def make_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = "SlideVoice — real slide playback QA"
    for label, title, body, notes, background in SLIDES:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*background)
        for top, height, text, size in (
            (0.65, 0.5, label, 18),
            (1.6, 1.6, title, 38),
            (3.8, 1.5, body, 25),
            (6.5, 0.5, "SLIDEVOICE  |  SYNTHETIC ENGLISH QA", 16),
        ):
            box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(height))
            box.text_frame.word_wrap = True
            box.text_frame.text = text
            for paragraph in box.text_frame.paragraphs:
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor(242, 247, 255)
        slide.notes_slide.notes_text_frame.text = notes
    deck.save(path)


def generate(root: Path) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    fixtures = root / "fixtures"
    fixtures.mkdir(exist_ok=True)
    slides, voice = fixtures / "slides.pptx", fixtures / "synthetic_reference.wav"
    make_deck(slides)
    write_tone(voice, 1, frequency=220)
    config = load_config(Path(__file__).resolve().parents[1] / "configs/qa_slides.yaml")
    job = create_job(slides, voice, root / "narrated", config)
    pipeline = Pipeline(job)
    project = pipeline.run()
    compatible = root / "slides_compatible.mp4"
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-i",
            str(job / "video.mp4"),
            "-map",
            "0:v:0",
            "-map",
            "0:a:0",
            "-map_metadata",
            "-1",
            "-map_chapters",
            "-1",
            "-c:v",
            "libx264",
            "-profile:v",
            "baseline",
            "-level:v",
            "3.1",
            "-preset",
            "fast",
            "-crf",
            "20",
            "-pix_fmt",
            "yuv420p",
            "-c:a",
            "copy",
            "-movflags",
            "+faststart",
            str(compatible),
        ],
        timeout=300,
    )
    fallback = root / "slides_fallback.webm"
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-i",
            str(compatible),
            "-c:v",
            "libvpx-vp9",
            "-deadline",
            "realtime",
            "-cpu-used",
            "8",
            "-crf",
            "35",
            "-b:v",
            "0",
            "-c:a",
            "libopus",
            "-b:a",
            "128k",
            str(fallback),
        ],
        timeout=300,
    )
    run_process(
        [
            "ffmpeg",
            "-v",
            "error",
            "-nostdin",
            "-y",
            "-i",
            str(compatible),
            "-map",
            "0:a:0",
            "-c:a",
            "pcm_s16le",
            str(job / "audio_check.wav"),
        ]
    )
    engine = WindowsSAPI({}, fixtures)
    engine.load_voice(
        VoiceProfile(
            ref_audio=voice,
            ref_text="unused",
            language="en",
            consent_confirmed=True,
        )
    )
    sides = []
    for word in ("Left", "Right"):
        path = fixtures / f"{word.lower()}.wav"
        engine.synthesize(f"{word} channel. This is the {word.lower()} headphone.", path)
        with wave.open(str(path), "rb") as reader:
            sides.append(reader.readframes(reader.getnframes()))
    stereo = bytearray()
    for index, samples in enumerate(sides):
        for i in range(0, len(samples), 2):
            stereo.extend(
                samples[i : i + 2] + b"\0\0" if index == 0 else b"\0\0" + samples[i : i + 2]
            )
        stereo.extend(b"\0" * 24000 * 4)
    with wave.open(str(root / "headphone_check.wav"), "wb") as writer:
        writer.setparams((2, 2, 24000, 0, "NONE", "not compressed"))
        writer.writeframes(stereo)
    probe = json.loads(
        run_process(
            [
                "ffprobe",
                "-v",
                "error",
                "-show_format",
                "-show_streams",
                "-of",
                "json",
                str(compatible),
            ]
        )
    )
    for video in (job / "video.mp4", compatible, fallback):
        run_process(
            [
                "ffmpeg",
                "-v",
                "error",
                "-xerror",
                "-nostdin",
                "-i",
                str(video),
                "-map",
                "0:v:0",
                "-map",
                "0:a:0",
                "-f",
                "null",
                "-",
            ],
            timeout=300,
        )
    write_json(
        root / "qa_checks.json",
        {
            "renderer": "Installed Microsoft PowerPoint exporting actual PPTX graphics",
            "slides": len(project.slides),
            "audio": measure_audio(compatible),
            "probe": probe,
            "full_decode_passed": ["narrated/video.mp4", compatible.name, fallback.name],
            "sha256": {p.name: file_hash(p) for p in (compatible, fallback)},
            "human_listening_verified": False,
            "browser_playback_verified": False,
            "note": "Decode and signal checks do not prove sound reaches physical headphones.",
        },
    )
    write_player(root)
    shutil.copyfile(job / "work/render/slide_001.png", root / "slide_preview.png")
    return job


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    generate(args.out.resolve())
